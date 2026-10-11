import pytest
from docx import Document

from cv_screen.extract import MockExtractor
from cv_screen.parsing import ParseError, find_cv_files, normalize_text, parse_file
from cv_screen.pipeline import screen_text
from tests.conftest import SAMPLES


def test_parses_pdf_sample():
    text = parse_file(SAMPLES / "cv01.pdf")
    assert "Customer Service Agent, Example Shop GmbH (fictional) - 03/2022 - present" in text
    assert "English: C1" in text
    assert "SYNTHETIC CV" in text


def test_parses_docx_sample_and_marks_bullets():
    text = parse_file(SAMPLES / "cv02.docx")
    assert "WORK EXPERIENCE" in text
    assert "- Resolved customer tickets by email and chat in Freshdesk" in text


def test_docx_tables_are_read(tmp_path):
    doc = Document()
    doc.add_paragraph("Languages")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "German"
    table.rows[0].cells[1].text = "C1"
    path = tmp_path / "table.docx"
    doc.save(path)
    assert "German | C1" in parse_file(path)


def test_txt_is_supported(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_text("Skills\r\nExcel, Zendesk\n\n\n\nEnd", encoding="utf-8")
    assert parse_file(path) == "Skills\nExcel, Zendesk\n\nEnd\n"


def test_unsupported_type_raises(tmp_path):
    path = tmp_path / "cv.odt"
    path.write_text("x")
    with pytest.raises(ParseError, match="unsupported"):
        parse_file(path)


def test_corrupt_pdf_raises_parse_error(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(ParseError):
        parse_file(path)


def test_empty_text_layer_raises(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    with pytest.raises(ParseError, match="no text"):
        parse_file(path)


def test_normalize_text_unifies_dashes_and_bullets():
    assert normalize_text("2019 – 2021\n• Item") == "2019 - 2021\n- Item\n"


def test_find_cv_files_filters_and_sorts(tmp_path):
    for name in ["b.pdf", "a.docx", "job.yaml", ".hidden.pdf", "README.md", "notes.txt"]:
        (tmp_path / name).write_text("x")
    assert [p.name for p in find_cv_files(tmp_path)] == ["a.docx", "b.pdf", "notes.txt"]


def test_invisible_format_characters_do_not_split_words(job, as_of):
    raw = "Skills\nZen\u00addesk, live\u200bchat\nLanguages\n\u202eGerman: C1\ufeff\n"
    text = normalize_text(raw)
    assert text == "Skills\nZendesk, livechat\nLanguages\nGerman: C1\n"
    cand = screen_text("c", "c.txt", text, job, MockExtractor(), as_of)
    assert {r.criterion_id: r.status for r in cand.results}["helpdesk_tool"] == "met"
