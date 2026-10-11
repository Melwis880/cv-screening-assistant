"""Changing protected attributes must not change the result.

The strongest form: the text handed to the extractor (mock or Claude) is byte-identical, so
nothing downstream can depend on those attributes. We also check the final match results.
"""

import re

import pytest
from docx import Document

from cv_screen.extract import ClaudeExtractor, MockExtractor
from cv_screen.parsing import parse_file
from cv_screen.pipeline import screen_text
from tests.conftest import SAMPLES
from tests.fakes import GOOD_PAYLOAD, FakeClient

VARIANTS = [
    {
        "name": "Erika Mustermann", "address": "12 Example Street, 10115 Sampletown",
        "email": "erika.mustermann@example.com", "phone": "+49 30 0000 0100",
        "dob": "14 March 1999", "nationality": "Fictionalian", "gender": "Female",
        "marital": "Single", "photo": "[Photo]", "pronoun": "She", "possessive": "her",
    },
    {
        "name": "Mohammed Al-Beispiel", "address": "1 Sample Road, 99999 Othertown",
        "email": "m.al-beispiel@example.org", "phone": "+44 20 0000 9999",
        "dob": "2 January 1961", "nationality": "Otherlandian", "gender": "Male",
        "marital": "Married, three children", "photo": "Photo", "pronoun": "He",
        "possessive": "his",
    },
    {
        "name": "Kim Ng", "address": "Flat 4, 7 Test Lane, Demo City",
        "email": "kim@example.net", "phone": "0151 0000 0000",
        "dob": "30 December 2003", "nationality": "Exampleland", "gender": "Non-binary",
        "marital": "Divorced", "photo": "", "pronoun": "She", "possessive": "her",
    },
    {   # names starting with a non-ASCII letter were not recognised as names before
        "name": "Özlem Yılmaz", "address": "Ringstraße 5, 50667 Beispielstadt",
        "email": "oezlem@example.com", "phone": "+49 221 0000 0000",
        "dob": "5 May 1988", "nationality": "Demotopian", "gender": "Female",
        "marital": "Married", "photo": "[Foto]", "pronoun": "She", "possessive": "her",
    },
    {   # nor were names with particles
        "name": "Joost van der Berg", "address": "Teststraat 9, 1011 AB Voorbeeld",
        "email": "joost.vdberg@example.nl", "phone": "+31 20 000 0000",
        "dob": "1 July 1975", "nationality": "Samplish", "gender": "Male",
        "marital": "Widowed", "photo": "Photo", "pronoun": "He", "possessive": "his",
    },
]

TEMPLATE = """{name}
{photo}
Address: {address}
Email: {email} | Phone: {phone}
Date of birth: {dob} | Nationality: {nationality} | Gender: {gender} | Marital status: {marital}
Profile
{first} is a calm support agent. {pronoun} has answered customer emails for years and {possessive}
team relies on {first}.
Experience
Customer Service Agent, Example Shop GmbH (fictional) - 03/2024 - present
- Answered customer emails and live chat in Zendesk
Cashier, Demo Market (fictional) - 01/2023 - 12/2023
Education
Vocational training as Retail Salesperson (IHK), Sampletown Vocational School, 2019 - 2022
Languages
German: C1
English: B2
Spanish: fluent
Skills
Zendesk, live chat, Excel
"""


def render(v):
    return TEMPLATE.format(first=v["name"].split()[0], **v)


def outcome(cand):
    return [(r.criterion_id, r.status, r.evidence) for r in cand.results], cand.counts


def test_redacted_text_and_results_identical_across_protected_attributes(job, as_of):
    runs = [screen_text("c", "c.txt", render(v), job, MockExtractor(), as_of) for v in VARIANTS]
    texts = {r.redacted_text for r in runs}
    assert len(texts) == 1, "text sent to the extractor differs between variants"
    assert len({str(outcome(r)) for r in runs}) == 1
    # and the result is non-trivial, so we are not comparing empty outputs
    statuses = {s for _, s, _ in outcome(runs[0])[0]}
    assert {"met", "partly", "unknown"} <= statuses


def test_claude_receives_identical_prompt_across_protected_attributes(job, as_of):
    prompts = []
    for v in VARIANTS:
        client = FakeClient(GOOD_PAYLOAD)
        screen_text("c", "c.txt", render(v), job, ClaudeExtractor(client=client), as_of)
        call = client.messages.calls[0]
        prompts.append((call["system"], call["messages"][0]["content"]))
    assert len(set(prompts)) == 1
    for v in VARIANTS:
        for key in ["name", "dob", "nationality", "email", "phone", "address"]:
            for token in re.split(r"[\s,]+", v[key]):
                # skip words that also occur in the job-relevant body ("Example Shop GmbH")
                if len(token) > 3 and not token.isdigit() and token not in TEMPLATE:
                    assert token not in prompts[0][1], (key, token)


@pytest.mark.parametrize("sample", ["cv01.pdf", "cv04.docx", "cv11.pdf", "cv14.docx"])
def test_sample_cv_invariance(sample, job, as_of):
    original = parse_file(SAMPLES / sample)
    swapped = original
    swaps = {
        r"Date of birth: [^|\n]+": "Date of birth: 1 January 1950 ",
        r"Nationality: [^|\n]+": "Nationality: Otherlandian ",
        r"Gender: [^|\n]+": "Gender: Diverse ",
        r"Address: [^\n]+": "Address: 99 Different Road, Faraway",
    }
    for pattern, repl in swaps.items():
        swapped = re.sub(pattern, repl, swapped)
    first_line = original.splitlines()[0]
    if first_line not in ("CURRICULUM VITAE", "[Photo]"):
        swapped = swapped.replace(first_line, "Zara Completely-Different", 1)
    swapped = re.sub(r"Name: [^\n]+", "Name: Zara Completely-Different", swapped)
    assert swapped != original

    a = screen_text("c", sample, original, job, MockExtractor(), as_of)
    b = screen_text("c", sample, swapped, job, MockExtractor(), as_of)
    assert a.redacted_text == b.redacted_text
    assert outcome(a) == outcome(b)


def test_docx_with_and_without_photo_and_header(tmp_path, job, as_of):
    def make(path, v, photo):
        doc = Document()
        if photo:
            from io import BytesIO

            from PIL import Image
            buf = BytesIO()
            Image.new("RGB", (20, 20), (1, 2, 3)).save(buf, format="PNG")
            doc.add_picture(buf)
        for line in render(v).splitlines():
            doc.add_paragraph(line)
        doc.save(path)

    make(tmp_path / "a.docx", VARIANTS[0], photo=True)
    make(tmp_path / "b.docx", VARIANTS[1], photo=False)
    a = screen_text("c", "a", parse_file(tmp_path / "a.docx"), job, MockExtractor(), as_of)
    b = screen_text("c", "b", parse_file(tmp_path / "b.docx"), job, MockExtractor(), as_of)
    assert a.redacted_text == b.redacted_text
    assert outcome(a) == outcome(b)


GERMAN_TEMPLATE = """Lebenslauf
{name}
Anschrift: {address}
Telefon: {phone} | E-Mail: {email}
Geburtsdatum: {dob} | Geburtsort: {birthplace} | Staatsangehörigkeit: {nationality}
Familienstand: {marital} | Geschlecht: {gender}
Profile
Calm support agent, {age}.
Experience
Customer Service Agent, Example Shop GmbH (fictional) - 03/2024 - present
- Answered customer emails and live chat in Zendesk
Languages
German: C1
English: B2
Skills
Zendesk, live chat, Excel
"""

GERMAN_VARIANTS = [
    {"name": "Erika Mustermann", "address": "Musterstraße 1, 10115 Berlin",
     "phone": "+49 30 0000 0100", "email": "erika@example.de", "dob": "14.03.1999",
     "birthplace": "Beispielstadt", "nationality": "deutsch", "marital": "ledig",
     "gender": "weiblich", "age": "27 Jahre alt"},
    {"name": "Ahmet Öztürk", "address": "Probeweg 7, 80331 München",
     "phone": "+49 89 0000 0200", "email": "a.oeztuerk@example.de", "dob": "02.01.1961",
     "birthplace": "Demokent", "nationality": "türkisch", "marital": "verheiratet, 3 Kinder",
     "gender": "männlich", "age": "65-year-old"},
]


def test_german_labels_and_free_text_age_are_removed(job, as_of):
    runs = [screen_text("c", "c.txt", GERMAN_TEMPLATE.format(**v), job, MockExtractor(), as_of)
            for v in GERMAN_VARIANTS]
    assert runs[0].redacted_text == runs[1].redacted_text
    assert outcome(runs[0]) == outcome(runs[1])
    for v in GERMAN_VARIANTS:
        for value in v.values():
            for token in re.split(r"[\s,]+", value):
                if len(token) > 3:
                    assert token not in runs[0].redacted_text, token
